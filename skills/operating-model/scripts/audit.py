"""Rule 6 unit audit (units-and-scale contract) over a finished operating-model workbook.

Usage:  python3 audit.py <workbook.xlsx> [--last-hist-col H]

Checks 1-5 and 7-11 are mechanical and reported as a findings table; check 6 prints the stored
and displayed value of `Total Revenue` in the last historical column for a one-line
comparison against the cited filing. Also prints the structural checklist items
(hidden audit columns, freeze panes, gridlines, title, hyperlinks on input cells).
Run `recalc.py` first so cached values exist.
"""
import re
import sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

import style as S

# ---- Lays when: the row-condition vocabulary (SKILL.md, "Lays when") --------------
MODES = {"", "ALWAYS_COMPLETE", "ALWAYS_DASHES_OK", "DATA_DEPENDENT"}
FLAGS = {
    "HAS_GROSS_PROFIT_ROW": ("COSTS", re.compile(r"^Gross profit( \(modeled\))?$")),
    "HAS_DNA_EMBEDDED": ("COSTS", "Memo: depreciation and amortization (embedded in the lines above)"),
    "HAS_SEGMENT_ROWS": ("COSTS", re.compile(r"^Total Costs and Expenses — .+")),
    "HAS_LOCATION_BLOCK": ("REVENUE", "Company-Operated EOP Balance"),
    "HAS_NET_SALES_ROW": ("REVENUE", "Net Sales"),
}
OPERATORS = {"NOT", "AND", "(", ")"}
SECTIONS = ["REVENUE", "COSTS", "CONSOLIDATED FINANCIALS", "DEPRECIATION SCHEDULE", "NET WORKING CAPITAL"]
SECTION_OF = {"revenue": "REVENUE", "costs": "COSTS", "consolidated-financials": "CONSOLIDATED FINANCIALS",
              "depreciation-schedule": "DEPRECIATION SCHEDULE", "net-working-capital": "NET WORKING CAPITAL"}


def layout_rows(text):
    """Rows of every table that carries a `Lays when` column: (labels, lays_when, line_no, label_only).
    labels = the backticked spans of the first cell (an `A` / `B` cell means any one of
    them); a label containing `<` is a template and is skipped by the presence check."""
    out, lines = [], text.split("\n")
    i, example = 0, False
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("#"):
            example = bool(re.search(r"worked (example|illustration)", ln, re.I))
        if not example and ln.startswith("|") and "Lays when" in ln and i + 1 < len(lines) and lines[i + 1].startswith("|---"):
            hdr = [c.strip() for c in ln.strip().strip("|").split("|")]
            li = next(j for j, h in enumerate(hdr) if h == "Lays when")
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if len(cells) > li:
                    labels = [l.strip() for span in re.findall(r"`([^`]+)`", cells[0]) for l in span.split(" / ")]
                    label_only = bool(re.search(r"label[- ]only|label row|sub-?header|block header|group header|\(header\)",
                                                " ".join(cells), re.I))
                    out.append((labels, cells[li].strip("`").strip(), i + 1, label_only))
                i += 1
            continue
        i += 1
    return out


def lays_tokens(value):
    return re.findall(r"[A-Z_]+|\(|\)", value)


def lays_valid(value):
    if value in MODES:
        return True
    toks = lays_tokens(value)
    return bool(toks) and all(t in FLAGS or t in OPERATORS for t in toks)


def lays_holds(value, flags):
    """True when the row is required on this sheet. `flags` maps flag name -> bool."""
    if value in ("", "ALWAYS_COMPLETE", "ALWAYS_DASHES_OK"):
        return True
    if value == "DATA_DEPENDENT":
        return False
    expr = " ".join({"NOT": "not", "AND": "and"}.get(t, t if t in "()" else str(bool(flags.get(t, False))))
                    for t in lays_tokens(value))
    return bool(eval(expr, {"__builtins__": {}}, {}))


if __name__ == "__main__":
    pass
path = sys.argv[1]
last_hist = "H"
if "--last-hist-col" in sys.argv:
    last_hist = sys.argv[sys.argv.index("--last-hist-col") + 1]
STAGES = sys.argv[sys.argv.index("--stages") + 1:] if "--stages" in sys.argv else []
LH = column_index_from_string(last_hist)

wb = load_workbook(path); ws = wb.active
wbv = load_workbook(path, data_only=True); wsv = wbv.active


def norm(fmt):
    return (fmt or "").replace('\\-', '"-"').replace('\\(', '(').replace('\\)', ')').replace('"-"?', '"-"').replace('\\ ', '')


# data region
last = ws.max_row
labels, tags = {}, {}
for r in range(S.FIRST_DATA_ROW, last + 1):
    labels[r] = ws.cell(r, 3).value; tags[r] = ws.cell(r, 4).value
    if labels[r] == "Notes":
        last = r - 1; break
max_col = max((c.column for row in ws.iter_rows(min_row=S.PERIOD_HEADER_ROW, max_row=S.PERIOD_HEADER_ROW) for c in row if c.value is not None), default=S.FIRST_DATA_COL)
cols = list(range(S.FIRST_DATA_COL, max_col + 1))

# money families (units contract, Rule 3): the expected format is derived from the
# historical input cells of each family, the same test the renderer applies
fam_vals = {}
for r in range(S.FIRST_DATA_ROW, last + 1):
    fam = S.money_family(tags[r])
    if not fam:
        continue
    for c in range(S.FIRST_DATA_COL, LH + 1):
        cell = ws.cell(r, c)
        if cell.fill.fgColor.rgb in (S.INPUT_FILL, S.TENK_FILL) and isinstance(cell.value, (int, float)):
            fam_vals.setdefault(fam, []).append(cell.value)
fam_fmt = {fam: S.money_format(fam, vals) for fam, vals in fam_vals.items()}

findings = []
for r in range(S.FIRST_DATA_ROW, last + 1):
    label, tag = labels[r], tags[r]
    cells = [ws.cell(r, c) for c in cols]
    isdata = any((c.value is not None and c.value != "-") for c in cells)
    # 1 tag completeness
    if isdata and not tag:
        findings.append((r, label, "1 tag", "no Col D tag", "tag"))
    # 2 no General format in the data region
    for c in cells:
        if c.value is not None and c.number_format == "General":
            findings.append((r, label, "2 General", f"{c.coordinate} General", "row format"))
    # 3 family consistency: format matches the tag
    if tag:
        exp = S.formats_for_tag(tag)
        fam = S.money_family(tag)
        if fam in fam_fmt:
            exp = (fam_fmt[fam],)  # derived precision: one format for the whole family
        if exp is None:
            findings.append((r, label, "3 family", f"unrecognized tag {tag}", "a tag listed in style.TAG_FORMATS"))
        else:
            for c in cells:
                if c.value is not None and norm(c.number_format) not in [norm(e) for e in exp]:
                    findings.append((r, label, "3 family", f"{c.coordinate} fmt {c.number_format!r} for tag {tag}", exp[0]))
    # 4 no scale constants in formulas
    for c in cells:
        v = c.value
        if isinstance(v, str) and v.startswith("=") and re.search(r"[*/]\s*1(000)+(?!\d)|[*/]\s*1e\d", v):
            findings.append((r, label, "4 scale const", f"{c.coordinate} {v}", "none"))
    # styling: real indents, sum bar never on C/D
    if isinstance(label, str) and label != label.lstrip():
        findings.append((r, label, "indent", "leading space in label", "real indent"))
    for cc in (3, 4):
        if ws.cell(r, cc).border.top.style:
            findings.append((r, label, "sumbar", "border on label col", "none"))
    # 11 header rows: a label-only row (no tag, no value or formula in any period cell) is a
    # header — bold + underlined label, period cells empty (styling rules, "Every header")
    if isinstance(label, str) and label and not tag and not isdata:
        if any(c.value == "-" for c in cells):
            findings.append((r, label, "11 header", "dashes in period cells", "empty cells"))
        fnt = ws.cell(r, 3).font
        if not (fnt.bold and fnt.underline):
            findings.append((r, label, "11 header", "label not bold + underlined", "bold + underlined"))


# 5 dimensional check on ratio rows of the shape =X/Y, =X/Y*365, =X/Y-1
def family(tag):
    if not tag: return None
    t = tag
    if t.startswith("(USD in"): return "money"
    if t.startswith("(USD per") or t.startswith("(USD millions per") or t.startswith("(USD thousands per"): return "perunit"
    if t.startswith("(%"): return "pct"
    if t in ("(days)", "(years)", "(x)"): return "ratio"
    if t.startswith("(sq ft in"): return "sqft"
    if t.startswith("(sq ft per"): return "sqftper"
    if re.match(r"\((people|users|pairs|units), in", t): return "units"
    if t in ("(locations)", "(stores)", "(restaurants)"): return "count"
    return "other"


def rowtag(ref):
    m = re.match(r"\$?([A-Z]+)\$?(\d+)", ref); rr = int(m.group(2))
    return tags.get(rr), labels.get(rr)


for r in range(S.FIRST_DATA_ROW, last + 1):
    v = ws.cell(r, LH).value
    if not (isinstance(v, str) and v.startswith("=")):
        continue
    m = re.match(r"^=([A-Z]+\d+)/([A-Z]+\d+)(\*365|-1)?$", v)
    if not m:
        continue
    nt, _ = rowtag(m.group(1)); dt, _ = rowtag(m.group(2)); tail = m.group(3); tag = tags[r]
    nf, df, tf = family(nt), family(dt), family(tag)
    ok = False
    if tail == "-1":
        ok = (tag == "(%)" and nt == dt)
    elif tail == "*365":
        ok = (tag == "(days)" and nf == "money" and df == "money")
    elif nf == "money" and df == "money":
        ok = tf in ("pct", "ratio")
    elif nf == df and nf in ("count", "units", "sqft"):
        ok = tf == "pct"
    elif nf == "money" and df in ("count", "units"):
        ok = tf == "perunit"
    elif nf == "money" and df == "sqft":
        ok = tf == "perunit"
    elif nf == "sqft" and df == "count":
        ok = tf == "sqftper"
    if not ok:
        findings.append((r, labels[r], "5 dimensional", f"{v}: {nt} / {dt} -> {tag}", "canceling tags"))

# 7-9 row presence, memo content, margin reads (SKILL.md Step 5.1)
section_of_row, cur = {}, None
for r in range(S.FIRST_DATA_ROW, last + 1):
    if isinstance(labels[r], str) and labels[r].strip() in SECTIONS:
        cur = labels[r].strip()
    section_of_row[r] = cur
rows_by_section = {sec: {} for sec in SECTIONS}      # label -> first row
rows_all = {sec: {} for sec in SECTIONS}             # label -> every row with that label
for r, sec in section_of_row.items():
    if sec and isinstance(labels[r], str):
        rows_by_section[sec].setdefault(labels[r].strip(), r)
        rows_all[sec].setdefault(labels[r].strip(), []).append(r)
hist_cols = list(range(S.FIRST_DATA_COL, LH + 1))


def all_dash(r):
    """True for a data row whose every historical cell is the text dash. A label-only row
    (sub-header, block header) has no cell at all and is not a data row."""
    vals = [ws.cell(r, c).value for c in hist_cols]
    return any(v == "-" for v in vals) and all(v in (None, "-") for v in vals)


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


flags = {}
for name, (sec, probe) in FLAGS.items():
    labs = rows_by_section[sec]
    flags[name] = any(probe.match(l) for l in labs) if hasattr(probe, "match") else (probe in labs)
if STAGES:
    print("stages:", " ".join(STAGES))
    print("flags:", {k: v for k, v in flags.items()})
    for stage in STAGES:
        text = open(stage).read()
        rel = stage.replace("\\", "/")
        top = next((k for k in SECTION_OF if f"/{k}/" in "/" + rel or rel.startswith(k)), None)
        sec = SECTION_OF.get(top)
        if sec is None:
            findings.append((0, stage, "7 presence", "stage file outside the section folders", "a stage file")); continue
        present = rows_by_section[sec]
        required = []
        def find(label):
            if label in present:
                return label
            return next((pl for pl in present if pl.startswith(label + " (must")), None)
        for labs, lays, ln, label_only in layout_rows(text):
            if not lays_valid(lays):
                findings.append((0, f"{rel}:{ln}", "7 presence", f"unknown Lays when {lays!r}", "a vocabulary value")); continue
            if not labs or any("<" in l or "…" in l for l in labs) or not lays_holds(lays, flags):
                continue
            hit = next((h for h in (find(l) for l in labs) if h), None)
            required.append(labs[0])
            if hit is None:
                findings.append((0, f"{rel}:{ln}", "7 presence", f"{' / '.join(labs)} not in {sec}", "row laid"))
            elif lays != "ALWAYS_DASHES_OK" and not label_only and all(all_dash(x) for x in rows_all[sec][hit]):
                findings.append((present[hit], hit, "7 presence", "all historical cells are dashes", "values or ALWAYS_DASHES_OK"))
        print(f"required rows from {rel}: {len(required)}")
rev_row = rows_by_section["REVENUE"].get("Total Revenue")
for r in range(S.FIRST_DATA_ROW, last + 1):
    lab = labels[r] if isinstance(labels[r], str) else ""
    m = re.search(r"\(must (=|≈) 0\)\s*$", lab)
    if not m:
        continue
    for c in hist_cols:
        v = num(wsv.cell(r, c).value)
        if v is None:
            continue
        if m.group(1) == "=":
            tol = 1.0
        else:
            tr = num(wsv.cell(rev_row, c).value) if rev_row else None
            tol = 0.005 * abs(tr) if tr else 1.0
        if abs(v) > tol:
            findings.append((r, lab, "8 memo", f"{ws.cell(r, c).coordinate} = {v:,.0f}", f"|v| ≤ {tol:,.0f}"))
costs = rows_by_section["COSTS"]
for margin, numer, denom in (("Gross margin", "Gross profit", "Total Revenue"), ("Operating margin", "Operating Income (modeled)", "Total Revenue")):
    mr = costs.get(margin); nr = costs.get(numer); dr = rows_by_section["REVENUE"].get(denom)
    if mr and nr and dr:
        mv, nv, dv = (num(wsv.cell(x, LH).value) for x in (mr, nr, dr))
        if None not in (mv, nv, dv) and dv and abs(mv - nv / dv) > 1e-6:
            findings.append((mr, margin, "9 margin", f"{mv:.4f} vs {nv / dv:.4f}", f"{numer} ÷ {denom}"))

# 10 vocabulary: Notes rows and cell comments are written for a reader who has not read the skill
BANNED = re.compile(r"§|\bSection \d|\bPath C?\d\b|\bCase [A-B]\d?\b|\bMode [AB]\b|\bTier \d\b|\bLayout [AB]\b|\bApproach \d\b|"
                    r"\bred driver|\bgreen input|\bgold (cell|override|fill)|\bcontract row|\bspine\b|\bhand-back|\bsub-skill|"
                    r"\brouter\b|\bDEFECT UPSTREAM|\bRE-ENTRY FAILURE|\bOUT OF ORDER", re.I)
notes_start = next((r for r in range(S.FIRST_DATA_ROW, ws.max_row + 1) if ws.cell(r, 3).value == "Notes"), None)
if notes_start:
    for r in range(notes_start + 1, ws.max_row + 1):
        v = ws.cell(r, 3).value
        if isinstance(v, str) and (m := BANNED.search(v)):
            findings.append((r, v[:40], "10 vocabulary", f"Notes: {m.group(0)!r}", "analyst words"))
for r in range(S.FIRST_DATA_ROW, last + 1):
    for c in range(3, max_col + 1):
        cm = ws.cell(r, c).comment
        if cm and (m := BANNED.search(cm.text or "")):
            findings.append((r, labels[r], "10 vocabulary", f"{ws.cell(r, c).coordinate} comment: {m.group(0)!r}", "analyst words"))

# 6 magnitude spot-check
for r in range(S.FIRST_DATA_ROW, last + 1):
    if labels[r] == "Total Revenue" and tags[r] and tags[r].startswith("(USD in"):
        stored = wsv.cell(r, LH).value
        scale = 1e6 if "millions" in tags[r] else 1e3
        dec = 1 if ".0" in (ws.cell(r, LH).number_format or "") else 0  # read the cell's own format
        shown = f"{stored/scale:,.{dec}f}" if isinstance(stored, (int, float)) else "n/a (recalc first)"
        print(f"6 magnitude: row {r} Total Revenue {last_hist}: stored {stored!r}; displayed {shown} {tags[r]} — compare to the filing")
        break

# structural checklist
dA = ws.column_dimensions["A"]
print("A/B hidden:", dA.hidden and (ws.column_dimensions["B"].hidden or dA.max >= 2),
      "| freeze:", ws.freeze_panes, "| gridlines:", ws.sheet_view.showGridLines, "| C1:", ws[S.TITLE_CELL].value)
noh = [ws.cell(r, c).coordinate for r in range(S.FIRST_DATA_ROW, last + 1) for c in cols
       if ws.cell(r, c).fill.fgColor.rgb in (S.INPUT_FILL, S.TENK_FILL) and not ws.cell(r, c).hyperlink]
print("input/10-K cells without hyperlink:", noh)
print(f"\nFINDINGS: {len(findings)}")
print("| Row | Col C label | Check | Observed | Expected |\n|---|---|---|---|---|")
for f_ in findings:
    print("| %s | %s | %s | %s | %s |" % f_)
