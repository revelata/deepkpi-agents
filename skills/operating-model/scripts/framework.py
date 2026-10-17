"""Compose-then-render workbook renderer for the operating-model skill.

How a build uses it (the mechanic SKILL.md "Build mechanics" describes):

  1. `configure(...)` once per build with the company title, output path and the
     Model timeline (historical and projection fiscal years, fiscal year-end dates).
  2. Each section is a Python module that returns a LIST OF ROW OBJECTS, nothing
     else. A row's cells are cell specs (`inp`, `tenk`, `f`, `drv`, `dash`, `zero`);
     a formula refers to another row by `{key}` placeholder, never by a sheet address,
     so a section may reference a row that another section lays down, regardless of
     which section is written first or where it sits on the sheet.
  3. `assemble({...})` places each section's rows beneath its ALL-CAPS header in the
     fixed scaffold order (REVENUE, COSTS, CONSOLIDATED FINANCIALS, DEPRECIATION
     SCHEDULE, NET WORKING CAPITAL). Build order and sheet order are independent.
  4. `render(rows, notes)` regenerates the WHOLE workbook from scratch: row numbers
     are assigned, every `{key}` is resolved to its address, styling is applied from
     style.py. Nothing is ever inserted into a saved workbook; re-run the build and
     every reference is re-resolved.

Module-level names (`HIST`, `PROJ`, `YEARS`, `COL`, `HCOL`, `PCOL`, `FYE`, `CFG`)
are containers mutated in place by `configure`, so `from framework import *` in a
section module sees the configured timeline whenever it is imported.
"""
import re
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

from style import *  # noqa: F401,F403  (colors, fonts, formats; re-exported to section modules)
import style as _style

SECTIONS = ("REVENUE", "COSTS", "CONSOLIDATED FINANCIALS", "DEPRECIATION SCHEDULE", "NET WORKING CAPITAL")

# ---- Model configuration (mutated in place by configure) -------------------
CFG = {"title": None, "out": None, "fye_note": None, "sheet": "Operating model"}
HIST, PROJ, YEARS = [], [], []
COL, FYE = {}, {}
HCOL, PCOL = [], []


def configure(*, title, out, hist, proj, fye, fye_note=None, sheet="Operating model"):
    """Set the model timeline and output once per build.

    title    : workbook title for C1, e.g. "Darden Restaurants, Inc. — operating model (deepKPI)"
    out      : output .xlsx path
    hist     : historical fiscal years, oldest first (the four N-3 .. N)
    proj     : projection fiscal years, oldest first (the five N+1 .. N+5)
    fye      : {fiscal year: datetime.date of that fiscal year's end}
    fye_note : short text for D3 describing the fiscal year-end convention
    """
    CFG.update(title=title, out=out, fye_note=fye_note, sheet=sheet)
    HIST[:] = list(hist); PROJ[:] = list(proj); YEARS[:] = HIST + PROJ
    COL.clear(); COL.update({y: get_column_letter(FIRST_DATA_COL + i) for i, y in enumerate(YEARS)})
    FYE.clear(); FYE.update(fye)
    HCOL[:] = [COL[y] for y in HIST]; PCOL[:] = [COL[y] for y in PROJ]


def last_col():
    return FIRST_DATA_COL + len(YEARS) - 1


def prev(y):
    """Column letter of the fiscal year before y (for BOP = prior EOP, prior-year bases)."""
    return COL[y - 1]


# ---- Rows and cell specs ---------------------------------------------------
class Row:
    """One sheet row. `key` makes it addressable from formulas as {key}.

    label/indent/bold/underline/italic/size style the Col C label; `tag` is the Col D
    units tag; `colA` is the hidden audit-trail text (deepKPI metric name); `fmt` the
    number format for every data cell; `sumbar` draws the single thin top border of a
    total row (a money row's decimals are derived per family by render(), see
    style.money_format, whatever `fmt` is passed); `cells` maps fiscal year -> cell spec; `dashes` fills unspecified years
    with a dash (data rows only: a row with no cells and no tag is a header, rendered bold +
    underlined with empty period cells whatever is passed); `hidden` places the row in an
    outline group, left expanded; `comment` is the
    Col C label comment (the operand recipe / methodology).
    """
    def __init__(self, key=None, label="", indent=0, bold=False, underline=False, italic=False,
                 tag=None, colA=None, fmt=None, sumbar=False, cells=None, comment=None,
                 size=None, dashes=False, hidden=False):
        self.key = key; self.label = label; self.indent = indent; self.bold = bold
        self.underline = underline; self.italic = italic; self.tag = tag; self.colA = colA
        self.fmt = fmt; self.sumbar = sumbar; self.cells = cells or {}; self.comment = comment
        self.size = size; self.dashes = dashes; self.hidden = hidden


def blank():
    return Row()


def section(label):
    """One of the five ALL-CAPS section headers (bold, underlined, size 16)."""
    return Row(label=label, bold=True, underline=True, size=SECTION_HEADER_SIZE)


def subheader(label, indent=1):
    """A block sub-header: bold + underlined; its line items sit one level deeper, its total at its level."""
    return Row(label=label, indent=indent, bold=True, underline=True)


def inp(value, url, note=None):
    """deepKPI input: pale-green fill, dark-blue hyperlinked value (provenance)."""
    return {"kind": "input", "v": value, "url": url, "note": note}


def tenk(value, url, note):
    """10-K text override of a deepKPI value: gold fill, hyperlinked, comment states why."""
    return {"kind": "tenk", "v": value, "url": url, "note": note}


def f(formula):
    """Live formula; other rows referenced as {key}, e.g. '=F{rev_total}-F{cost_total}'."""
    return {"kind": "formula", "v": formula}


def drv(value, note=None, url=None):
    """Forward driver: red font, dashed box. `value` may be a number or a {key} formula.
    `url` hyperlinks the cell to the filing a guided seed was read from."""
    return {"kind": "driver", "v": value, "note": note, "url": url}


def dash():
    return {"kind": "dash"}


def zero():
    return {"kind": "zero"}


def drivers(key, hist_cells, note=None):
    """The default driver seed (SKILL.md, "Build mechanics", rung 3): N+1 = the most
    recent numeric value among N, N-1, N-2 (0 when none is numeric); later years =
    the N+1 cell, so the assumption is flat and an override of N+1 moves the line.

    `hist_cells` are the row's historical cell specs (usually live ratios); the row's
    own key is needed so the seed can address them. The guard is ISNUMBER, never
    IFERROR, so a genuine error in a historical cell still surfaces: dashes are the
    text "-" and blanks are blank, and ISNUMBER is false for both. A guided seed or
    an issuer-normalized seed (rungs 1 and 2) is written with drivers_fixed();
    a recurring episodic cost line with drivers_runrate().
    """
    c = dict(hist_cells)
    n, n1, n2 = (f"{COL[y]}{{{key}}}" for y in (HIST[-1], HIST[-2], HIST[-3]))
    c[PROJ[0]] = drv(f"=IF(ISNUMBER({n}),{n},IF(ISNUMBER({n1}),{n1},IF(ISNUMBER({n2}),{n2},0)))", note)
    for y in PROJ[1:]:
        c[y] = drv(f"={COL[PROJ[0]]}{{{key}}}", note)
    return c


def drivers_runrate(key, hist_cells, note=None):
    """The run-rate seed for a recurring episodic cost line (costs/episodicLines.md):
    N+1 = median of the numeric values among N-2, N-1, N (0 when none are numeric);
    later years = the N+1 cell. The median is the spike-resistant run-rate: with one
    distorted year among three it always lands on a normal year, so the builder never
    has to decide which year is the outlier. COUNT-guarded: 0 when no seed year is numeric.
    """
    c = dict(hist_cells)
    a, b = COL[HIST[-3]], COL[HIST[-1]]
    rng = f"{a}{{{key}}}:{b}{{{key}}}"
    c[PROJ[0]] = drv(f"=IF(COUNT({rng})=0,0,MEDIAN({rng}))", note)
    for y in PROJ[1:]:
        c[y] = drv(f"={COL[PROJ[0]]}{{{key}}}", note)
    return c


def drivers_fixed(key, hist_cells, value, note=None, url=None):
    """A driver seeded at a stated value instead of the latest-year rule: the guided
    midpoint (rung 1, `url` = the earnings-release exhibit), an issuer-normalized
    latest-year figure (rung 2), a reference to the latest undistorted year such as
    "=G{key}" (rung 4), 0 for a one-off episodic line, or a formula such as an
    expected remaining program cost divided by N+1 revenue (costs/episodicLines.md).
    `value` is a number or a {key} formula string. Later years = the N+1 cell, so an
    override of N+1 moves the line.
    """
    c = dict(hist_cells)
    c[PROJ[0]] = drv(value, note, url)
    for y in PROJ[1:]:
        c[y] = drv(f"={COL[PROJ[0]]}{{{key}}}", note)
    return c


# ---- Assemble in scaffold order --------------------------------------------
def assemble(parts):
    """Place each section's rows beneath its header, in the fixed scaffold order.

    `parts` maps a section header (one of SECTIONS) to that section's row list. A
    section not yet built gets its header and two spacer rows, which is exactly the
    Step 2 scaffold when every section is missing. Row lists must not include the
    section header row themselves.
    """
    unknown = set(parts) - set(SECTIONS)
    if unknown:
        raise ValueError(f"unknown section(s): {sorted(unknown)}; use {SECTIONS}")
    rows = []
    for name in SECTIONS:
        body = list(parts.get(name) or [])
        rows.append(section(name))
        rows += body if body else [blank(), blank()]
        if body and not (body[-1].label == "" and not body[-1].cells):
            rows.append(blank())
    return rows


def scaffold():
    """Step 2: the empty, styled workbook with the five section headers."""
    return render(assemble({}), [])


# ---- Render -----------------------------------------------------------------
def _font(**kw):
    kw.setdefault("name", FONT_NAME); kw.setdefault("size", FONT_SIZE)
    return Font(**kw)


def render(rows, notes, out=None):
    """Regenerate the workbook from `rows`. Returns {key: row number}."""
    if not YEARS:
        raise RuntimeError("call configure(...) before render()")
    out = out or CFG["out"]
    rownum = {}
    r = FIRST_DATA_ROW
    for row in rows:
        row.n = r
        if row.key:
            if row.key in rownum:
                raise ValueError(f"duplicate row key {row.key!r}")
            rownum[row.key] = r
        r += 1
    last_row = r - 1

    unresolved = {}
    def sub(s, row):
        def repl(m):
            k = m.group(1)
            if k not in rownum:
                unresolved.setdefault(k, set()).add(row.label or row.key or f"row {row.n}")
                return "#REF"
            return str(rownum[k])
        return re.sub(r"\{(\w+)\}", repl, s)

    thin = Side(style=SUM_BAR_SIDE); dashed = Side(style=DRIVER_BORDER_SIDE)
    wb = Workbook(); ws = wb.active; ws.title = CFG["sheet"]
    ws.sheet_view.showGridLines = SHOW_GRIDLINES
    for c in HIDE_COLUMNS:
        ws.column_dimensions[c].hidden = True
    ws.column_dimensions["C"].width = COL_WIDTH_LABEL; ws.column_dimensions["D"].width = COL_WIDTH_UNITS
    for y in YEARS:
        ws.column_dimensions[COL[y]].width = COL_WIDTH_PERIOD
    ws[TITLE_CELL] = CFG["title"]; ws[TITLE_CELL].font = _font(bold=True); ws.merge_cells("C1:D1")
    for y in YEARS:
        c = ws[f"{COL[y]}{PERIOD_HEADER_ROW}"]; c.value = FYE[y]
        c.number_format = FMT_PERIOD_HEADER.format(year=y)
        c.font = _font(bold=True, underline="single"); c.alignment = Alignment(horizontal="right")
    if CFG["fye_note"]:
        d = ws[f"D{PERIOD_HEADER_ROW}"]; d.value = CFG["fye_note"]; d.font = _font(italic=True)
    for col in range(FIRST_DATA_COL, last_col() + 1):
        ws.column_dimensions[get_column_letter(col)].outlineLevel = 1
    ws.freeze_panes = FREEZE_PANES

    # Indent is set from the row's kind, not by the stage (styling rules, "Indentation"):
    # a header keeps the indent the stage gave it; the rows beneath it sit one level
    # deeper; a total (bold) returns to the header's level; a memo row sits at the items'
    # level; a trend or ratio row sits one level deeper than the line above it.
    hdr = 0; line = 0
    for row in rows:
        if not row.label and not row.cells:
            continue
        if row.size == SECTION_HEADER_SIZE:
            row.indent = 0; hdr = 0; line = 0; continue
        is_header = not row.cells and not row.tag  # label-only: no cells and no tag
        if is_header:
            hdr = row.indent if row.indent else hdr + 1
            row.indent = hdr; line = hdr; continue
        lab = row.label or ""
        is_memo = lab.startswith("Memo:") or lab.endswith("(memo)") or "(reported)" in lab
        if row.bold:
            row.indent = hdr; line = hdr
        elif is_memo:
            row.indent = hdr + 1
        elif row.italic:
            row.indent = line + 1
        else:
            row.indent = hdr + 1; line = row.indent

    # Money precision is derived, not chosen (units contract, Rule 3): one format per money
    # family, from the stored historical inputs — one decimal when any input is not a whole
    # number of the display unit, else whole. Overrides whatever format the stage passed.
    fam_inputs = {}
    for row in rows:
        fam = _style.money_family(row.tag)
        if fam and row.cells:
            for y in HIST:
                spec = row.cells.get(y)
                if spec and spec.get("kind") in ("input", "tenk"):
                    fam_inputs.setdefault(fam, []).append(spec.get("v"))
    fam_fmt = {fam: _style.money_format(fam, vals) for fam, vals in fam_inputs.items()}
    for row in rows:
        fam = _style.money_family(row.tag)
        if fam in fam_fmt:
            row.fmt = fam_fmt[fam]

    for row in rows:
        n = row.n
        ws.cell(row=n, column=2, value=n).font = _font()
        if row.colA:
            ws.cell(row=n, column=1, value=row.colA).font = _font(color=SOURCE_LABEL_FONT)
        c = ws.cell(row=n, column=3, value=row.label or None)
        is_hdr = bool(row.label) and not row.cells and not row.tag  # styling rules: every header is bold + underlined
        c.font = _font(size=row.size or FONT_SIZE, bold=row.bold or is_hdr,
                       underline="single" if (row.underline or is_hdr) else None, italic=row.italic)
        c.alignment = Alignment(horizontal="left", vertical="bottom", indent=row.indent)
        if row.comment:
            c.comment = Comment(row.comment, AUTHOR_MODEL)
        if row.tag:
            d = ws.cell(row=n, column=4, value=row.tag); d.font = _font(italic=True)
        is_data = bool(row.cells) or (row.dashes and bool(row.tag))  # a header never carries dashes
        for y in YEARS:
            cell = ws[f"{COL[y]}{n}"]
            spec = row.cells.get(y) if row.cells else None
            if spec is None:
                if not is_data:
                    continue
                spec = {"kind": "dash"}
            kind = spec["kind"]
            cell.number_format = row.fmt or FMT_COUNT
            cell.alignment = Alignment(horizontal="right")
            font = dict(bold=row.bold, italic=row.italic)
            if kind == "dash":
                cell.value = "-"; cell.font = _font(**font)
            elif kind == "zero":
                cell.value = 0; cell.font = _font(**font)
            elif kind == "input":
                cell.value = spec["v"]; cell.hyperlink = spec["url"]
                cell.fill = PatternFill("solid", fgColor=INPUT_FILL); cell.font = _font(color=LINK_FONT, **font)
                if spec.get("note"):
                    cell.comment = Comment(spec["note"], AUTHOR_INPUT)
            elif kind == "tenk":
                cell.value = spec["v"]; cell.hyperlink = spec["url"]
                cell.fill = PatternFill("solid", fgColor=TENK_FILL); cell.font = _font(color=LINK_FONT, **font)
                cell.comment = Comment(spec["note"], AUTHOR_TENK)
            elif kind == "formula":
                cell.value = sub(spec["v"], row); cell.font = _font(**font)
            elif kind == "driver":
                v = spec["v"]; cell.value = sub(v, row) if isinstance(v, str) else v
                if spec.get("url"):
                    cell.hyperlink = spec["url"]
                cell.font = _font(color=DRIVER_FONT, **font)
                cell.border = Border(left=dashed, right=dashed, top=dashed, bottom=dashed)
                if spec.get("note"):
                    cell.comment = Comment(spec["note"], AUTHOR_DRIVER)
            else:
                raise ValueError(f"unknown cell kind {kind!r} in row {row.label!r}")
            if row.sumbar:
                b = cell.border
                cell.border = Border(left=b.left, right=b.right, bottom=b.bottom, top=thin)

    if unresolved:
        detail = "; ".join(f"{{{k}}} referenced by {sorted(v)}" for k, v in sorted(unresolved.items()))
        raise KeyError(f"unresolved row key(s): {detail}. The section that lays that row has not been assembled.")

    # supporting rows -> outline groups, left expanded (no row is hidden or collapsed by default)
    i = 0
    while i < len(rows):
        if rows[i].hidden:
            j = i
            while j + 1 < len(rows) and rows[j + 1].hidden:
                j += 1
            ws.row_dimensions.group(rows[i].n, rows[j].n, hidden=False, outline_level=1)
            i = j + 1
        else:
            i += 1

    # Notes block beneath the last row
    nr = last_row + 3
    ws.cell(row=nr, column=3, value="Notes").font = _font(bold=True, underline="single")
    for i, note in enumerate(notes, 1):
        ws.cell(row=nr + i, column=3, value=f"{i}. {note}").font = _font()
    wb.save(out)
    return rownum
